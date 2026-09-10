"""Lectura de los archivos de planta y carga a la base.

Muestras de humedad : .xlsx del Wagner, o el .csv del informe (bloque MUESTRAS).
Programa de secado   : .xlsx del PLC (preferido, precision completa) o el .pdf
                       impreso (respaldo: trunca cada celda a 3 caracteres y
                       pierde el decimal en el 73% de los valores).
"""
import hashlib
import math
import os
import re
import statistics as st
import zlib

import db

PMIN, PMAX = 24, 26            # tablas por paquete
CAP = {"p12": 24, "p36": 27}   # paquetes por camara
NOMBRE_GRUPO = {"p12": "Camaras 1-2", "p36": "Camaras 3-6"}


class ErrorIngesta(Exception):
    pass


# --------------------------------------------------------------- utilidades
def sha256(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace("%", "").replace(" ", "")
    if not t:
        return None
    if "," in t and "." not in t:
        t = t.replace(",", ".")
    else:
        t = t.replace(",", "")
    try:
        f = float(t)
    except ValueError:
        return None
    return f if math.isfinite(f) else None


def camara_de_nombre(nombre):
    """'archive_Camara Secado 3_2026...' -> 3   (tolera 'Camera')"""
    m = re.search(r"cam[ae]ra\s*(?:de\s*)?secado\s*(\d+)", nombre, re.I)
    return int(m.group(1)) if m else None


# ----------------------------------------------- reconocimiento de la camara
def detectar(n):
    """Un paquete lleva de 24 a 26 tablas; de ahi sale el rango de paquetes
    posible, y con el la camara compatible."""
    kmin, kmax = math.ceil(n / PMAX), n // PMIN
    d = {
        "n": n, "kmin": kmin, "kmax": kmax, "reparte": kmin <= kmax,
        "c12": kmin <= CAP["p12"], "c36": kmin <= CAP["p36"],
    }
    d["ok"] = d["c36"]
    d["grupo"] = "p36"          # la regla solo puede *probar* 3-6
    d["ambiguo"] = d["c12"]
    return d


def k_preferido(d, grupo):
    cap = CAP[grupo]
    if not d["reparte"]:
        return min(cap, max(1, d["kmin"]))
    ks = [k for k in range(d["kmin"], d["kmax"] + 1) if k <= cap]
    return ks[-1] if ks else d["kmin"]


def repartir(n, k):
    """k paquetes lo mas parejos posible; los grandes se distribuyen a lo largo."""
    base, r = divmod(n, k)
    return [base + (1 if (i + 1) * r // k > i * r // k else 0) for i in range(k)]


# ----------------------------------------------------- muestras de humedad
def _muestras_desde_filas(filas):
    """Acepta el bloque 'Muestra;Valor' en pares, o una columna de humedad."""
    pares, cab = {}, None
    for f in filas:
        celdas = [("" if c is None else str(c).strip()) for c in f]
        if cab is None and any(re.fullmatch(r"muestra", c, re.I) for c in celdas):
            cab = True
            continue
        if cab:
            for k in range(0, len(celdas) - 1, 2):
                i, v = num(celdas[k]), num(celdas[k + 1])
                if i and v and i >= 1 and 0 < v < 80 and float(i).is_integer():
                    pares[int(i)] = v
    if len(pares) >= 2:
        return [pares[i] for i in sorted(pares)]

    # respaldo: la columna con mas numeros plausibles
    ancho = max((len(f) for f in filas), default=0)
    mejor, mejor_n = None, 0
    for c in range(ancho):
        vals = [num(f[c]) for f in filas if c < len(f)]
        vals = [v for v in vals if v is not None and 0 < v < 80]
        if len(vals) > mejor_n:
            mejor, mejor_n = vals, len(vals)
    if mejor_n < 2:
        raise ErrorIngesta("No se encontro una columna de humedad con datos numericos.")
    return mejor


def fecha_iso(s):
    """El informe escribe 'dd-mm-aaaa H:MM:SS'; en la base todo va en ISO."""
    if not s:
        return None
    s = " ".join(str(s).split())
    m = re.match(r"(\d{1,2})-(\d{1,2})-(\d{4})\s+(\d{1,2}):(\d{2})(?::(\d{2}))?", s)
    if m:
        d, mo, y, h, mi, se = m.groups()
        return "%s-%02d-%02d %02d:%s:%s" % (y, int(mo), int(d), int(h), mi, se or "00")
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})(?::(\d{2}))?", s)
    if m:
        y, mo, d, h, mi, se = m.groups()
        return "%s-%s-%s %s:%s:%s" % (y, mo, d, h, mi, se or "00")
    return None


def _cabecera_informe(texto):
    meta = {}
    lineas = texto.split("\n")
    for i, l in enumerate(lineas):
        if l.startswith("Ciclo;Muestras") and i + 1 < len(lineas):
            v = lineas[i + 1].split(";")
            meta["ciclo"] = int(num(v[0])) if num(v[0]) else None
            meta["especie"] = v[2].strip() if len(v) > 2 else None
        if l.startswith("Fecha Inicio"):
            # 'Fecha Inicio;26-08-2026 ; 8:12:11;;Fecha Termino;01-09-2026 ; 9:26:22'
            par = re.findall(r"(\d{1,2}-\d{1,2}-\d{4})\s*;\s*(\d{1,2}:\d{2}(?::\d{2})?)", l)
            if len(par) >= 1:
                meta["inicio_informe"] = fecha_iso("%s %s" % par[0])
            if len(par) >= 2:
                meta["fin_informe"] = fecha_iso("%s %s" % par[1])
        if l.startswith("Resultado:"):
            meta["resultado"] = l.split(";")[1].strip() if ";" in l else None
    return meta


def leer_muestras(ruta):
    ext = os.path.splitext(ruta)[1].lower()
    meta = {}
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        filas = []
        for hoja in wb.worksheets:
            filas += [list(r) for r in hoja.iter_rows(values_only=True)]
        texto = "\n".join(";".join("" if c is None else str(c) for c in f) for f in filas)
        meta = _cabecera_informe(texto)
        valores = _muestras_desde_filas(filas)
        fmt = "xlsx"
    else:
        with open(ruta, "rb") as f:
            crudo = f.read()
        try:
            texto = crudo.decode("utf-8")
        except UnicodeDecodeError:
            texto = crudo.decode("latin-1")
        meta = _cabecera_informe(texto)
        filas = [l.split(";") for l in texto.split("\n")]
        valores = _muestras_desde_filas(filas)
        fmt = "csv"

    if len(valores) < 2:
        raise ErrorIngesta("El archivo de muestras no trae mediciones.")
    meta["valores"] = valores
    meta["formato"] = fmt
    if not meta.get("ciclo"):
        m = re.search(r"(\d{4})", os.path.basename(ruta))
        meta["ciclo"] = int(m.group(1)) if m else None
    return meta


# ------------------------------------------------------- programa de secado
def _proceso_xlsx(ruta):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    next(it, None)                       # cabecera
    filas = []
    for r in it:
        if not r or r[0] is None:
            continue
        ts = str(r[0])[:19].replace("T", " ")
        if not re.match(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", ts):
            continue
        filas.append((ts, [num(v) for v in r[1:20]]))
    return filas, True                   # precision completa


def _proceso_pdf(ruta):
    """Extractor propio: el PDF trae el texto en un CID desplazado 29 posiciones
    y las columnas solo se distinguen por su coordenada X."""
    with open(ruta, "rb") as f:
        d = f.read()
    objs = {int(m.group(1)): m.end() for m in re.finditer(rb"[\n](\d+) 0 obj", d)}

    def crudo(n):
        s = objs[n]
        return d[s:d.find(b"endobj", s)]

    def flujo(n):
        o = crudo(n)
        st_ = o.find(b"stream") + 6
        b = o[st_:].lstrip(b"\r\n")
        return zlib.decompress(b[:b.rfind(b"endstream")])

    paginas = []
    for n in sorted(objs):
        o = crudo(n)
        if b"/Type/Page/" in o:
            paginas.append(int(re.search(rb"/Contents (\d+) 0 R", o).group(1)))

    BS, LP, RP = 92, 40, 41
    ESC = {ord("n"): 10, ord("r"): 13, ord("t"): 9, ord("b"): 8, ord("f"): 12,
           LP: LP, RP: RP, BS: BS}
    TM = re.compile(rb"1 0 0 1 ([0-9.]+) ([0-9.]+) Tm[\n][(]")

    def celdas(t):
        for m in TM.finditer(t):
            i, hondo, out = m.end(), 1, bytearray()
            while i < len(t):
                c = t[i]
                if c == BS:
                    sig = t[i + 1]
                    if sig in ESC:
                        out.append(ESC[sig]); i += 2
                    elif 48 <= sig <= 55:
                        j, oc = i + 1, ""
                        while j < len(t) and len(oc) < 3 and 48 <= t[j] <= 55:
                            oc += chr(t[j]); j += 1
                        out.append(int(oc, 8) & 0xFF); i = j
                    elif sig == 10:
                        i += 2
                    else:
                        out.append(sig); i += 2
                elif c == LP:
                    hondo += 1; out.append(c); i += 1
                elif c == RP:
                    hondo -= 1
                    if hondo == 0:
                        i += 1; break
                    out.append(c); i += 1
                else:
                    out.append(c); i += 1
            s = "".join(chr(out[k] * 256 + out[k + 1] + 29) for k in range(0, len(out) - 1, 2))
            yield float(m.group(1)), float(m.group(2)), s

    # Las celdas van centradas en su columna, asi que la x se corre segun el
    # largo del texto: la fecha cae en 58.93 o en 61.99 segun tenga la hora uno
    # o dos digitos. Por eso la fecha se busca por patron y cada valor se asigna
    # a la columna mas cercana, no a una x fija.
    colx = [170 + 19.7368 * i + 2.22 for i in range(19)]
    PASO = 19.7368 / 2
    TS = re.compile(r"\d{2}-\d{2}-\d{4} \d{1,2}:\d{2}:\d{2}")
    filas = {}
    for cn in paginas:
        por_y = {}
        for x, y, s in celdas(flujo(cn)):
            if y < 60:
                continue
            por_y.setdefault(round(y, 1), {})[x] = s
        for cc in por_y.items():
            cc = cc[1]
            ts = next((s for x, s in cc.items() if x < 170 and TS.fullmatch(s.strip())), None)
            if not ts:
                continue
            vals = [None] * 19
            for x, s in cc.items():
                if x < 170:
                    continue
                i = min(range(19), key=lambda k: abs(colx[k] - x))
                if abs(colx[i] - x) <= PASO:
                    vals[i] = num(s)
            if sum(v is not None for v in vals) < 4:   # cabeceras y pies de pagina
                continue
            filas[ts.strip()] = vals

    def clave(t):
        dd, mm, yy = t[:2], t[3:5], t[6:10]
        return "%s-%s-%s %s" % (yy, mm, dd, t[11:].rjust(8, "0"))

    return [(clave(t), filas[t]) for t in sorted(filas, key=clave)], False


def leer_proceso(ruta, nombre=None):
    """`nombre` es el nombre original del archivo: al subirlo por la web la ruta
    es un temporal y la camara solo se puede leer del nombre que trajo."""
    ext = os.path.splitext(nombre or ruta)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        filas, exacto = _proceso_xlsx(ruta)
        fmt = "xlsx"
    elif ext == ".pdf":
        filas, exacto = _proceso_pdf(ruta)
        fmt = "pdf"
    else:
        raise ErrorIngesta("El programa de secado debe ser .xlsx del PLC o .pdf impreso.")
    if not filas:
        raise ErrorIngesta("No se pudieron leer registros del programa de secado.")
    return {"filas": filas, "formato": fmt, "precision": exacto,
            "camara": camara_de_nombre(nombre or os.path.basename(ruta))}


# ------------------------------------------------------------------- carga
def puntuar(cx, camara, ciclo, valores, objetivo, lo, hi, umbral):
    """Clasifica las tablas contra la banda [lo, hi] y deja el veredicto en ciclo.

    La banda de aceptacion se guarda como dos limites, no como objetivo +- tol:
    no tiene por que estar centrada en la humedad que persigue el secado."""
    if hi < lo:
        lo, hi = hi, lo
    bajas = sum(1 for x in valores if x < lo)
    altas = sum(1 for x in valores if x > hi)
    n = len(valores)
    pct = 100.0 * (n - bajas - altas) / n
    cx.execute("""UPDATE ciclo SET n_bajas=?, n_altas=?, n_rango=?, resultado=?,
                  crit_objetivo=?, crit_lo=?, crit_hi=?, crit_umbral=?,
                  actualizado_en=datetime('now','localtime')
                  WHERE camara=? AND ciclo=?""",
               (bajas, altas, n - bajas - altas,
                "Aprobado" if pct >= umbral else "Rechazado",
                objetivo, lo, hi, umbral, camara, ciclo))
    return {"n_bajas": bajas, "n_altas": altas, "n_rango": n - bajas - altas, "pct_rango": pct}


def recriterio(cx, camara, ciclo, objetivo=None, lo=None, hi=None, umbral=None):
    """Vuelve a puntuar un ciclo ya cargado con otra banda, sin releer archivos.
    Los parametros en None conservan lo que ya tenia el ciclo."""
    c = cx.execute("SELECT crit_objetivo, crit_lo, crit_hi, crit_umbral FROM ciclo"
                   " WHERE camara=? AND ciclo=?", (camara, ciclo)).fetchone()
    if c is None:
        raise ErrorIngesta("La camara %s ciclo %s no esta en la base." % (camara, ciclo))
    v = [r["mc"] for r in db.muestras(cx, camara, ciclo)]
    if not v:
        raise ErrorIngesta("El ciclo %s no tiene muestras que puntuar." % ciclo)
    r = puntuar(cx, camara, ciclo, v,
                c["crit_objetivo"] if objetivo is None else objetivo,
                c["crit_lo"] if lo is None else lo,
                c["crit_hi"] if hi is None else hi,
                c["crit_umbral"] if umbral is None else umbral)
    cx.commit()
    return r


def guardar(cx, camara, ciclo, muestras_ruta=None, proceso_ruta=None, nombres=None,
            objetivo=9.0, lo=8.0, hi=12.0, umbral=85.0):
    """Carga uno o ambos archivos de un ciclo. Devuelve un informe de la carga."""
    avisos, det, nombres = [], None, nombres or {}
    cx.execute("INSERT OR IGNORE INTO ciclo (camara, ciclo) VALUES (?,?)", (camara, ciclo))

    # ---- muestras
    if muestras_ruta:
        m = leer_muestras(muestras_ruta)
        v = m["valores"]
        det = detectar(len(v))
        k = k_preferido(det, det["grupo"])
        tam, off, filas = repartir(len(v), k), 0, []
        for p, t in enumerate(tam):
            for j in range(t):
                filas.append((camara, ciclo, off + j + 1, v[off + j], p, j))
            off += t
        cx.execute("DELETE FROM muestra WHERE camara=? AND ciclo=?", (camara, ciclo))
        cx.executemany("INSERT INTO muestra (camara,ciclo,n_muestra,mc,paquete,posicion)"
                       " VALUES (?,?,?,?,?,?)", filas)
        cx.execute("""UPDATE ciclo SET n_muestras=?, mc_medio=?, mc_sd=?, mc_min=?, mc_max=?,
                      k_paquetes=?, grupo_camara=?,
                      ambiguo=?, especie=COALESCE(?,especie),
                      inicio_informe=COALESCE(?,inicio_informe), fin_informe=COALESCE(?,fin_informe),
                      actualizado_en=datetime('now','localtime')
                      WHERE camara=? AND ciclo=?""",
                   (len(v), st.mean(v), st.stdev(v) if len(v) > 1 else 0, min(v), max(v),
                    k, det["grupo"], 1 if det["ambiguo"] else 0, m.get("especie"),
                    m.get("inicio_informe"), m.get("fin_informe"), camara, ciclo))
        puntuar(cx, camara, ciclo, v, objetivo, lo, hi, umbral)
        _archivo(cx, camara, ciclo, "muestras", muestras_ruta, m["formato"], True,
                 nombres.get("muestras"))

        if not det["ok"]:
            avisos.append(("alerta",
                "%d muestras necesitan al menos %d paquetes y la camara mas grande admite %d: "
                "el archivo trae mas de una carga o el Wagner duplico los datos."
                % (det["n"], det["kmin"], CAP["p36"])))
        elif det["ambiguo"]:
            avisos.append(("aviso",
                "%d muestras dan entre %d y %d paquetes, que caben tanto en camaras 1-2 como en 3-6: "
                "el total no basta para distinguirlas." % (det["n"], det["kmin"], det["kmax"])))
        else:
            avisos.append(("ok",
                "%d muestras dan entre %d y %d paquetes; mas de %d solo cabe en %s."
                % (det["n"], det["kmin"], det["kmax"], CAP["p12"], NOMBRE_GRUPO["p36"])))

    # ---- programa de secado
    if proceso_ruta:
        p = leer_proceso(proceso_ruta, nombres.get("proceso"))
        if p["camara"] and p["camara"] != camara:
            avisos.append(("aviso", "El nombre del archivo dice camara %d y se esta cargando "
                                    "en la camara %d." % (p["camara"], camara)))
        cx.execute("DELETE FROM proceso WHERE camara=? AND ciclo=?", (camara, ciclo))
        cols = ",".join(db.COLS_PROCESO)
        hueco = ",".join("?" * (3 + len(db.COLS_PROCESO)))
        cx.executemany("INSERT OR REPLACE INTO proceso (camara,ciclo,ts,%s) VALUES (%s)" % (cols, hueco),
                       [(camara, ciclo, ts) + tuple(vals[:19] + [None] * (19 - len(vals)))
                        for ts, vals in p["filas"]])
        cx.execute("""UPDATE ciclo SET inicio_log=?, fin_log=?,
                      actualizado_en=datetime('now','localtime') WHERE camara=? AND ciclo=?""",
                   (p["filas"][0][0], p["filas"][-1][0], camara, ciclo))
        _archivo(cx, camara, ciclo, "proceso", proceso_ruta, p["formato"], p["precision"],
                 nombres.get("proceso"))
        if not p["precision"]:
            avisos.append(("aviso",
                "El PDF trunca cada celda a 3 caracteres: los valores de dos digitos pierden el "
                "decimal. Si existe el .xlsx del PLC, cargalo en su lugar."))
        else:
            avisos.append(("ok", "%d registros del PLC con precision completa." % len(p["filas"])))

    for nivel, msg in avisos:
        db.registrar(cx, camara, ciclo, nivel, msg)
    cx.commit()
    return {"avisos": avisos, "deteccion": det}


def _archivo(cx, camara, ciclo, tipo, ruta, formato, precision, nombre=None):
    destino = os.path.join(db.RUTA_ARCHIVOS, "%d_%d_%s%s"
                           % (camara, ciclo, tipo,
                              os.path.splitext(nombre or ruta)[1].lower()))
    if os.path.abspath(ruta) != os.path.abspath(destino):
        with open(ruta, "rb") as a, open(destino, "wb") as b:
            b.write(a.read())
    cx.execute("""INSERT OR REPLACE INTO archivo
                  (camara,ciclo,tipo,nombre,sha256,bytes,formato,precision_decimal,ruta,cargado_en)
                  VALUES (?,?,?,?,?,?,?,?,?,datetime('now','localtime'))""",
               (camara, ciclo, tipo, nombre or os.path.basename(ruta), sha256(destino),
                os.path.getsize(destino), formato, 1 if precision else 0, destino))
